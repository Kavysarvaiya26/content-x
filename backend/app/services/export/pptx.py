from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

from app.schemas import PresentationOutline


def render_pptx(outline: PresentationOutline) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for slide in outline.slides:
        layout = prs.slide_layouts[0 if slide.layout == "title" else 1]
        s = prs.slides.add_slide(layout)
        if s.shapes.title:
            s.shapes.title.text = slide.title or outline.title
        body = None
        for shape in s.shapes:
            if shape.has_text_frame and shape != s.shapes.title:
                body = shape
                break
        if body and body.has_text_frame:
            tf = body.text_frame
            tf.clear()
            bullets = list(slide.bullets)
            if slide.layout == "two_column":
                bullets = bullets + slide.right_column_bullets
            if not bullets:
                tf.text = outline.subtitle or outline.footer
            for i, bullet in enumerate(bullets):
                paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                paragraph.text = bullet.text
                paragraph.level = min(bullet.level, 4)
                paragraph.font.size = Pt(18)
        if s.has_notes_slide:
            s.notes_slide.notes_text_frame.text = slide.speaker_notes
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
